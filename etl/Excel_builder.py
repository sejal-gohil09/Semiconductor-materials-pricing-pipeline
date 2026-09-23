import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

def build_consolidated_excel(clean_df, extrap_df, survey_df, output_path="reports/Semiconductor_Materials_Market_Insights.xlsx"):
    wb = openpyxl.Workbook()
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    # Sheet 1: Consolidate Prices across files
    ws1 = wb.active
    ws1.title = "Consolidated Material Prices"
    for r in dataframe_to_rows(clean_df, index=False, header=True):
        ws1.append(r)
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Sheet 2: Extrapolation Models
    ws2 = wb.create_sheet(title="Price Extrapolation Forecast")
    for r in dataframe_to_rows(extrap_df, index=False, header=True):
        ws2.append(r)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Sheet 3: Market Sentiment
    if not survey_df.empty:
        ws3 = wb.create_sheet(title="Survey Market Sentiment")
        for r in dataframe_to_rows(survey_df, index=False, header=True):
            ws3.append(r)
        for cell in ws3[1]:
            cell.font = header_font
            cell.fill = header_fill

    wb.save(output_path)
    print(f"[EXCEL BUILDER] Successfully generated executive report at '{output_path}'.")
